# -*- coding: utf-8 -*-
"""Convierte Informe2.md en Informe2.docx: carta, margenes 2,5 cm, Calibri 11,
cuerpo justificado, encabezados en negrita sin numerar, tablas markdown con fila
de encabezado sombreada y figuras insertadas desde ![pie](ruta).

El bloque de cabecera (lineas antes del primer "## ") va una por linea, sin
negrita, salvo la ultima, que es el titulo.

USO
    python informe_a_word.py
"""
import io
import os
import re

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(AQUI, "Informe2.md")
SALIDA = os.path.join(AQUI, "Informe2.docx")
ACENTO = "2E5090"
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)


def corridas(p, texto, tam=11, negrita=False, cursiva=False, color=None):
    for trozo in re.split(r"(\*\*.+?\*\*|\*.+?\*)", texto):
        if not trozo:
            continue
        neg, cur, t = negrita, cursiva, trozo
        if trozo.startswith("**"):
            neg, t = True, trozo[2:-2]
        elif trozo.startswith("*"):
            cur, t = True, trozo[1:-1]
        r = p.add_run(t)
        r.font.size = Pt(tam)
        r.bold = neg or None
        r.italic = cur or None
        if color is not None:
            r.font.color.rgb = color


def parrafo(d, texto, tam=11, justificar=True, despues=6, negrita=False, cursiva=False, centrar=False):
    p = d.add_paragraph()
    corridas(p, texto, tam, negrita, cursiva)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centrar else (
        WD_ALIGN_PARAGRAPH.JUSTIFY if justificar else WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.space_after = Pt(despues)
    return p


def sombrear(celda, hex_color):
    pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    pr.append(shd)


def tabla(d, filas):
    celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
    celdas = [c for c in celdas if not all(re.fullmatch(r":?-+:?", x) for x in c)]
    enc, cuerpo = celdas[0], celdas[1:]
    t = d.add_table(rows=1, cols=len(enc))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, c in enumerate(enc):
        cel = t.rows[0].cells[k]; cel.text = ""
        corridas(cel.paragraphs[0], c, 9.5, negrita=True, color=BLANCO)
        cel.paragraphs[0].paragraph_format.space_after = Pt(0)
        sombrear(cel, ACENTO)
    tr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); tr.append(h)
    for fila in cuerpo:
        nueva = t.add_row().cells
        for k, c in enumerate(fila[:len(enc)]):
            nueva[k].text = ""
            corridas(nueva[k].paragraphs[0], c, 9.5)
            nueva[k].paragraphs[0].paragraph_format.space_after = Pt(0)
            if k > 0 and re.fullmatch(r"[−\-+]?\d[\d.,]*( ?\(.*\))?|—", c):
                nueva[k].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    lineas = io.open(MD, encoding="utf-8").read().split("\n")
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(2.5))
    est = d.styles["Normal"]
    est.font.name = "Calibri"; est.font.size = Pt(11)
    est.paragraph_format.line_spacing = 1.15

    primer = next(i for i, l in enumerate(lineas) if l.startswith("## "))
    cab = [l for l in lineas[:primer] if l.strip()]
    for l in cab[:-1]:
        parrafo(d, l, tam=11, justificar=False, despues=0, centrar=True)
    parrafo(d, cab[-1], tam=13, negrita=True, despues=14, centrar=True)

    i = primer
    buf_tabla = []
    while i < len(lineas):
        l = lineas[i]
        t = l.strip()
        if t.startswith("|"):
            buf_tabla.append(t); i += 1; continue
        if buf_tabla:
            tabla(d, buf_tabla); buf_tabla = []
        if not t:
            i += 1; continue
        if t.startswith("## "):
            p = parrafo(d, t[3:], tam=12, negrita=True, justificar=False, despues=4)
            p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
        elif t.startswith("### "):
            p = parrafo(d, t[4:], tam=11, negrita=True, justificar=False, despues=3)
            p.paragraph_format.space_before = Pt(6); p.paragraph_format.keep_with_next = True
        elif t.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.*)\)", t)
            pie, ruta = m.group(1), os.path.join(AQUI, m.group(2))
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(ruta, width=Inches(6.3))
            p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2)
            parrafo(d, pie, tam=9.5, cursiva=True, despues=10)
        elif t.startswith("Tabla ") or t.startswith("Table "):
            p = parrafo(d, t, tam=9.5, cursiva=True, justificar=False, despues=3, centrar=True)
            p.paragraph_format.keep_with_next = True
        else:
            parrafo(d, t)
        i += 1
    if buf_tabla:
        tabla(d, buf_tabla)
    d.save(SALIDA)
    v = docx.Document(SALIDA)
    print("escrito: %s | parrafos %d | tablas %d | imagenes %d" % (
        SALIDA, len(v.paragraphs), len(v.tables), len([r for r in v.part.rels.values() if "image" in r.reltype])))


if __name__ == "__main__":
    main()
